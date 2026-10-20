"""Bounded local format adapters; all semantic rendering lives in tables.py."""

from __future__ import annotations

import csv
from datetime import date, datetime, time, timedelta
import io
import re
import zipfile
from xml.etree import ElementTree as ET

from tele_bot.knowledge.tables import (
    StructuredTable, TableCell, TableLimits, TableMerge, TableReadOptions, TableRow, build_table,
)


def checked_archive(data: bytes, limits: TableLimits) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        limits.check(len(archive.infolist()), 10000, "压缩包条目数")
        limits.check(sum(info.file_size for info in archive.infolist()), limits.max_unpacked_bytes, "解压后字节数")
        if len(archive.namelist()) != len(set(archive.namelist())):
            raise ValueError("文档压缩包含重复条目")
        return archive
    except zipfile.BadZipFile as exc:
        raise ValueError("invalid Office document archive") from exc


def safe_xml(data: bytes):
    if re.search(br"<!\s*(?:DOCTYPE|ENTITY)", data.replace(b"\x00", b""), re.I):
        raise ValueError("Office XML entities/DOCTYPE are not supported")
    try:
        return ET.fromstring(data)
    except ET.ParseError as exc:
        raise ValueError("invalid Office XML") from exc


def decode_csv(data: bytes, options: TableReadOptions) -> str:
    for encoding in ((options.encoding,) if options.encoding else ("utf-8-sig", "utf-16" if data.startswith((b'\xff\xfe', b'\xfe\xff')) else "gb18030")):
        try:
            return data.decode(encoding)
        except (UnicodeError, LookupError):
            continue
    raise ValueError("CSV 编码不可读，请通过 table-options.encoding 明确指定编码")


def read_csv(name: str, data: bytes, options: TableReadOptions, limits: TableLimits) -> tuple[StructuredTable, ...]:
    if options.sheets or options.cell_range:
        raise ValueError("CSV 不支持工作表/单元格范围选择，请拆分 CSV 文件")
    text = decode_csv(data, options)
    delimiter = options.delimiter
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows, cells_count = [], 0
    try:
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
        for record, values in enumerate(reader, 1):
            limits.check(record, limits.max_rows, "逻辑记录数")
            limits.check(len(values), limits.max_columns, "列数")
            cells_count += len(values)
            limits.check(cells_count, limits.max_cells, "单元格数")
            for value in values:
                limits.check(len(value), limits.max_cell_chars, "单元格字符数")
            # Empty records preserve logical numbering, without creating fake facts.
            if values:
                rows.append(TableRow(record, tuple(TableCell(record, col, value if value != "" else None)
                                                  for col, value in enumerate(values, 1))))
    except csv.Error as exc:
        raise ValueError(f"CSV 解析失败：{exc}；请检查分隔符、引号或缩小单元格") from exc
    if not rows:
        raise ValueError("CSV contains no records")
    if options.header_rows > len(rows):
        raise ValueError("header_rows exceeds available CSV records")
    return (build_table(source="csv", identity=name, title=name, rows=tuple(rows),
                        header_rows=options.header_rows,
                        region=f"records:{rows[0].position}-{rows[-1].position}"),)


def _value(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def read_xlsx(name: str, data: bytes, options: TableReadOptions, limits: TableLimits) -> tuple[StructuredTable, ...]:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries, get_column_letter
    except ImportError as exc:
        raise RuntimeError("XLSX parsing requires openpyxl==3.1.5") from exc
    # Preflight BEFORE openpyxl materializes rows/merge ranges. A tiny ZIP with
    # malicious dimensions must not cause a huge in-memory worksheet allocation.
    total_cells, total_area, total_merge_area, sheet_count = 0, 0, 0, 0
    with checked_archive(data, limits) as archive:
        for path in archive.namelist():
            if not path.endswith((".xml", ".rels")):
                continue
            xml = safe_xml(archive.read(path))
            if xml.tag.rsplit("}", 1)[-1] == "worksheet":
                sheet_count += 1
                limits.check(sheet_count, limits.max_sheets, "工作表数")
                max_row, max_column = 1, 1
                for node in xml.iter():
                    tag = node.tag.rsplit("}", 1)[-1]
                    if tag in {"dimension", "mergeCell"}:
                        ref = node.get("ref", "A1")
                        left, top, right, bottom = range_boundaries(ref)
                        if any(v is None for v in (left, top, right, bottom)):
                            raise ValueError("invalid XLSX dimensions")
                        limits.check(bottom, limits.max_rows, "源行号")
                        limits.check(right, limits.max_columns, "源列号")
                        limits.check(bottom * right, limits.max_cells, "工作表区域单元格数")
                        max_row, max_column = max(max_row, bottom), max(max_column, right)
                        if tag == "mergeCell":
                            total_merge_area += (bottom-top+1) * (right-left+1)
                            limits.check(total_merge_area, limits.max_cells, "合并区域总单元格数")
                    if tag == "c":
                        total_cells += 1
                        limits.check(total_cells, limits.max_cells, "总单元格数")
                        ref = node.get("r", "A1")
                        left, top, _, _ = range_boundaries(ref)
                        if left is None or top is None:
                            raise ValueError("invalid XLSX cell coordinate")
                        limits.check(top, limits.max_rows, "源行号")
                        limits.check(left, limits.max_columns, "源列号")
                        max_row, max_column = max(max_row, top), max(max_column, left)
                total_area += max_row * max_column
                limits.check(total_area, limits.max_cells, "工作簿区域总单元格数")
            for node in xml.iter():
                if node.tag.rsplit("}", 1)[-1] in {"t", "f", "v"}:
                    limits.check(len(node.text or ""), limits.max_cell_chars, "单元格字符数")
    try:
        formulas = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
        cached = load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    except Exception as exc:
        raise ValueError("invalid or unreadable XLSX document") from exc
    tables = []
    try:
        selected = options.sheets or tuple(formulas.sheetnames)
        if any(sheet not in formulas.sheetnames for sheet in selected):
            raise ValueError("选定工作表不存在，请核对 sheets 配置")
        for sheet_number, sheet_name in enumerate(formulas.sheetnames, 1):
            if sheet_name not in selected:
                continue
            ws, values = formulas[sheet_name], cached[sheet_name]
            left, top, right, bottom = (range_boundaries(options.cell_range) if options.cell_range
                                        else (1, 1, ws.max_column, ws.max_row))
            if any(v is None for v in (left, top, right, bottom)) or min(left, top) < 1 or right < left or bottom < top:
                raise ValueError("cell_range must be a bounded A1 range")
            limits.check(bottom, limits.max_rows, "选定源行号")
            limits.check(right, limits.max_columns, "选定源列号")
            limits.check((right-left+1) * (bottom-top+1), limits.max_cells, "选定区域单元格数")
            merges = tuple(TableMerge(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges)
            for merge in merges:
                intersects = merge.first_row <= bottom and merge.last_row >= top and merge.first_column <= right and merge.last_column >= left
                if intersects and not (top <= merge.first_row <= merge.last_row <= bottom and left <= merge.first_column <= merge.last_column <= right):
                    raise ValueError("选定区域切断合并单元格，请包含完整合并范围")
            grid = {}
            for row in ws.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right):
                for cell in row:
                    if cell.value is not None:
                        is_formula = cell.data_type == "f"
                        value = _value(values.cell(cell.row, cell.column).value if is_formula else cell.value)
                        formula = (cell.value if isinstance(cell.value, str) else getattr(cell.value, "text", None)) if is_formula else None
                        if is_formula and not isinstance(formula, str):
                            raise ValueError(f"{sheet_name}!{cell.coordinate}：公式类型不可解释，请将其转换为静态值后导入")
                        limits.check(len(value or "") + len(formula or ""), limits.max_cell_chars, "单元格字符数")
                        grid[cell.row, cell.column] = TableCell(cell.row, cell.column, value, formula, cell.number_format)
            # Blank row/column separators are boundaries, not records or headers.
            occupied = set(grid)
            for merge in merges:
                if (merge.first_row, merge.first_column) in grid:
                    occupied.update((r, c) for r in range(merge.first_row, merge.last_row + 1)
                                    for c in range(merge.first_column, merge.last_column + 1))
            row_bands = [(top, bottom)] if options.cell_range else _bands(sorted({r for r, _ in occupied}))
            for row_band in row_bands:
                columns = sorted({c for r, c in occupied if row_band[0] <= r <= row_band[1]})
                column_bands = [(left, right)] if options.cell_range else _bands(columns)
                for col_band in column_bands:
                    first, last = row_band
                    start, stop = col_band
                    region = f"{get_column_letter(start)}{first}:{get_column_letter(stop)}{last}"
                    rows = tuple(TableRow(r, tuple(grid.get((r,c), TableCell(r,c,None))
                                                   for c in range(start, stop+1))) for r in range(first,last+1))
                    if options.header_rows > len(rows):
                        raise ValueError(f"{sheet_name}!{region}: header_rows exceeds region rows")
                    local_merges = tuple(m for m in merges if first <= m.first_row <= m.last_row <= last
                                         and start <= m.first_column <= m.last_column <= stop)
                    tables.append(build_table(source="xlsx", identity=f"{sheet_number}:{sheet_name}:{region}",
                                              title=f"{name} / {sheet_name} / 区域{len(tables)+1}", rows=rows,
                                              header_rows=options.header_rows, merges=local_merges, region=region,
                                              section_path=(sheet_name,)))
    finally:
        formulas.close()
        cached.close()
    if not tables:
        raise ValueError("XLSX contains no selected extractable cells")
    return tuple(tables)


def _bands(positions):
    if not positions:
        return []
    result, start, last = [], positions[0], positions[0]
    for value in positions[1:]:
        if value != last + 1:
            result.append((start, last))
            start = value
        last = value
    return result + [(start, last)]
