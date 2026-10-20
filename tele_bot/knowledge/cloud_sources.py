"""Explicit cloud selections -> complete snapshots -> the shared text/table path."""

from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import datetime, timezone
import json
import re
from urllib.parse import urlparse, parse_qs, urlencode

from openpyxl.utils.cell import range_boundaries, get_column_letter

from tele_bot.knowledge.cloud_fields import (
    FieldValueError,
    field_text,
    SUPPORTED_FIELD_TYPES,
)
from tele_bot.knowledge.feishu_docs import CloudAPIError, FeishuDocsClient, resource_id
from tele_bot.knowledge.ingestion import DocumentParser, ParsedDocument, sha256_text
from tele_bot.knowledge.tables import (
    TableCell,
    TableRow,
    TableMerge,
    TableColumn,
    StructuredTable,
    build_table,
    render_table,
)


def canonical(value):
    return json.dumps(
        value,
        sort_keys=True,
        ensure_ascii=False,
        separators=(",", ":"),
        allow_nan=False,
    )


@dataclass(frozen=True)
class CloudResource:
    kind: str
    token: str
    host: str
    entry_url: str
    query_json: str = "{}"


def parse_cloud_url(url):
    parsed = urlparse(url.strip())
    host = (parsed.hostname or "").lower()
    if (
        parsed.scheme != "https"
        or parsed.username
        or parsed.password
        or parsed.port not in {None, 443}
        or not any(
            host == d or host.endswith("." + d) for d in ("feishu.cn", "larksuite.com")
        )
    ):
        raise ValueError("仅支持官方 Feishu/Lark HTTPS 云文档链接")
    match = re.fullmatch(
        r"/(docx|wiki|sheets|base|docs)/([A-Za-z0-9_-]{1,256})/?", parsed.path
    )
    if not match or match[1] == "docs":
        raise ValueError("不支持该云文档类型（旧版 /docs/ 需先转换为新版 docx）")
    query = parse_qs(parsed.query, keep_blank_values=True)
    if any(len(v) != 1 for v in query.values()):
        raise ValueError("重复的云文档链接选择参数")
    return CloudResource(
        match[1],
        resource_id(match[2]),
        host,
        url.strip(),
        canonical({k: v[0] for k, v in query.items()}),
    )


def normalize_selection(resource, selection):
    if not isinstance(selection, dict):
        raise ValueError("--options 必须是 JSON 对象")
    selection = json.loads(canonical(selection))
    query = json.loads(resource.query_json)
    if resource.kind == "wiki":
        return selection  # Validate after resolving the actual resource type.
    allowed = {
        "docx": {"header_rows", "embeds"},
        "sheets": {"header_rows", "sheet_id", "range"},
        "base": {"table_id", "field_ids", "filter", "view_id"},
    }[resource.kind]
    if set(selection) - allowed:
        raise ValueError(
            "不支持的云来源选择参数：" + ", ".join(sorted(set(selection) - allowed))
        )
    if resource.kind in {"sheets", "docx"}:
        selection.setdefault("header_rows", 1)
        if (
            type(selection["header_rows"]) is not int
            or not 0 <= selection["header_rows"] <= 100
        ):
            raise ValueError("header_rows 必须为 0–100 的整数")
    if resource.kind == "sheets":
        selection.setdefault("sheet_id", query.get("sheet", ""))
        selection.setdefault("range", query.get("range", ""))
        if not selection.get("sheet_id") or not selection.get("range"):
            raise ValueError(
                'Sheets 必须明确 --options {"sheet_id":"...","range":"A1:F200"}；不会自动读取其他工作表'
            )
        resource_id(selection["sheet_id"])
        if not isinstance(selection["range"], str) or not re.fullmatch(
            r"[A-Z]+[1-9]\d*:[A-Z]+[1-9]\d*", selection["range"]
        ):
            raise ValueError("range 必须是完整 A1 区域，例如 A1:F200")
    elif resource.kind == "base":
        selection.setdefault("table_id", query.get("table", ""))
        if query.get("view"):
            selection.setdefault("view_id", query["view"])
        if not selection.get("table_id") or not selection.get("field_ids"):
            raise ValueError(
                "Base 必须明确选择 table_id 和非空 field_ids；可用 /kb inspect-doc 查看元数据"
            )
        resource_id(selection["table_id"])
        ids = selection["field_ids"]
        if not isinstance(ids, list) or any(not isinstance(i, str) for i in ids):
            raise ValueError("field_ids 必须是字段身份列表")
        selection["field_ids"] = sorted(set(resource_id(i) for i in ids))
        if "view_id" in selection:
            resource_id(selection["view_id"])
        if "filter" in selection and not isinstance(selection["filter"], dict):
            raise ValueError("filter 必须是显式 API 过滤对象")
    else:
        selection.setdefault("embeds", {})
        if not isinstance(selection["embeds"], dict) or len(selection["embeds"]) > 32:
            raise ValueError("embeds 必须为至多 32 个 block_id 的选择配置")
    return selection


@dataclass(frozen=True)
class CloudSnapshot:
    resource: CloudResource
    selection_json: str
    source_ref: str
    parsed: ParsedDocument
    revision: str | None
    read_at: str
    children: tuple[CloudSnapshot, ...] = ()
    references_json: str = "[]"


class CloudSourceConnector:
    def __init__(
        self, client: FeishuDocsClient, parser: DocumentParser, *, batch_rows=200
    ):
        self.client, self.parser = client, parser
        self.limits = parser.table_limits
        self.batch_rows = batch_rows

    def resolve(self, url):
        resource = parse_cloud_url(url)
        is_lark = resource.host.endswith("larksuite.com")
        if is_lark != (self.client.api_base_url == "https://open.larksuite.com"):
            raise ValueError("来源地域与配置的 Docs API 域不一致")
        if resource.kind == "wiki":
            node = self.client.get(
                "/open-apis/wiki/v2/spaces/get_node", token=resource.token
            ).get("node", {})
            kind = {"docx": "docx", "sheet": "sheets", "bitable": "base"}.get(
                node.get("obj_type")
            )
            if not kind:
                raise ValueError(f"Wiki 实际资源类型暂不支持：{node.get('obj_type')}")
            resource = replace(
                resource, kind=kind, token=resource_id(node.get("obj_token"))
            )
        return resource

    def source_ref(self, resource, selection):
        return "cloud:" + sha256_text(
            canonical(
                [
                    self.client.api_base_url,
                    self.client.app_id,
                    resource.kind,
                    resource.token,
                    normalize_selection(resource, selection),
                ]
            )
        )

    def actual_url(self, resource, selection):
        suffix = ""
        if resource.kind == "sheets":
            suffix = "?" + urlencode(
                {"sheet": selection["sheet_id"], "range": selection["range"]}
            )
        elif resource.kind == "base":
            suffix = "?" + urlencode({"table": selection["table_id"]})
        return f"https://{resource.host}/{resource.kind}/{resource.token}{suffix}"

    def inspect(self, url):
        resource = self.resolve(url)
        token = resource.token
        if resource.kind == "sheets":
            data = self.client.get(
                f"/open-apis/sheets/v3/spreadsheets/{token}/sheets/query"
            )
            return {
                "kind": "sheets",
                "resource_token": token,
                "sheets": data.get("sheets", []),
                "requires": "sheet_id + bounded range",
            }
        if resource.kind == "base":
            tables = self.client.pages(
                f"/open-apis/bitable/v1/apps/{token}/tables", params={"page_size": 100}
            )
            self.limits.check(len(tables), self.limits.max_sheets, "Base 数据表数")
            definitions = []
            for table in tables:
                tid = resource_id(table["table_id"])
                fields = self.client.pages(
                    f"/open-apis/bitable/v1/apps/{token}/tables/{tid}/fields",
                    params={"page_size": 100},
                )
                definitions.append(
                    {
                        "table_id": tid,
                        "name": table.get("name"),
                        "fields": [
                            {k: f.get(k) for k in ("field_id", "field_name", "type")}
                            for f in fields
                        ],
                    }
                )
            return {
                "kind": "base",
                "tables": definitions,
                "requires": "table_id + field_ids; view is not an ACL boundary",
            }
        snapshot = self.load(resource, {})
        return {
            "kind": "docx",
            "title": snapshot.parsed.name,
            "references": json.loads(snapshot.references_json),
            "missing": snapshot.parsed.missing_content,
        }

    def load(self, resource, selection):
        with self.client.read_scope():
            try:
                return self._load(resource, selection)
            except (KeyError, TypeError, AttributeError, IndexError) as exc:
                raise CloudAPIError(
                    "invalid_response", "云端结构不符合已支持契约，未发布不完整快照"
                ) from exc

    def _load(self, resource, selection):
        selection = normalize_selection(resource, selection)
        result = getattr(self, "_" + resource.kind)(resource, selection)
        pieces, tables, missing, revision, children, references, title = result
        text, sections = self.parser._assemble(pieces)
        self.limits.check(
            len(text) + sum(len(c.parsed.text) for c in children),
            self.limits.max_output_chars,
            "云来源及嵌入输出字符数",
        )
        # Empty complete tables/records still produce a source section, allowing
        # confirmed deletion snapshots to replace the previous record set.
        parsed = ParsedDocument(
            title,
            text,
            sha256_text(text),
            "text/plain",
            sections,
            tuple(tables),
            parser_profile=canonical({"cloud_renderer": "v1", "selection": selection}),
            missing_content=tuple(missing),
        )
        return CloudSnapshot(
            resource,
            canonical(selection),
            self.source_ref(resource, selection),
            parsed,
            str(revision) if revision is not None else None,
            datetime.now(timezone.utc).isoformat(),
            tuple(children),
            canonical(references),
        )

    def _table_parts(self, table):
        return [
            (p.text, p.path, p.locator, "table")
            for p in render_table(
                table, max_chars=self.parser.table_chunk_chars, limits=self.limits
            )
        ]

    def _sheets(self, resource, selection):
        token, sid = resource.token, selection["sheet_id"]
        meta = self.client.get(
            f"/open-apis/sheets/v3/spreadsheets/{token}/sheets/query"
        )
        sheet = next(
            (s for s in meta.get("sheets", []) if s.get("sheet_id") == sid), None
        )
        if sheet is None:
            raise CloudAPIError("not_found", "选定工作表不存在")
        sheet = self.client.get(
            f"/open-apis/sheets/v3/spreadsheets/{token}/sheets/{sid}"
        ).get("sheet")
        if not isinstance(sheet, dict) or sheet.get("sheet_id") != sid:
            raise CloudAPIError("invalid_response", "Sheets 缺少选定工作表元数据")
        if sheet.get("resource_type", "sheet").lower() not in {"sheet", "spreadsheet"}:
            raise ValueError("选定工作表不是可读取单元格表格")
        left, top, right, bottom = range_boundaries(selection["range"])
        if right < left or bottom < top:
            raise ValueError("range 起止顺序无效")
        self.limits.check(bottom, self.limits.max_rows, "Sheets 源行号")
        self.limits.check(right, self.limits.max_columns, "Sheets 源列号")
        self.limits.check(
            (bottom - top + 1) * (right - left + 1),
            self.limits.max_cells,
            "Sheets 区域单元格数",
        )
        merges = []
        for m in sheet.get("merges", []):
            merge = TableMerge(
                m["start_row_index"] + 1,
                m["start_column_index"] + 1,
                m["end_row_index"],
                m["end_column_index"],
            )
            if (
                merge.first_row <= bottom
                and merge.last_row >= top
                and merge.first_column <= right
                and merge.last_column >= left
            ):
                if not (
                    top <= merge.first_row <= merge.last_row <= bottom
                    and left <= merge.first_column <= merge.last_column <= right
                ):
                    raise ValueError("Sheets 选定区域切断合并单元格")
                merges.append(merge)
        rows, revision = [], None
        for first in range(top, bottom + 1, self.batch_rows):
            last = min(bottom, first + self.batch_rows - 1)
            area = f"{sid}!{get_column_letter(left)}{first}:{get_column_letter(right)}{last}"
            data = self.client.get(
                f"/open-apis/sheets/v2/spreadsheets/{token}/values/{area}",
                valueRenderOption="ToString",
                dateTimeRenderOption="FormattedString",
            )
            current = data.get("revision")
            if current is None:
                raise CloudAPIError(
                    "invalid_response", "Sheets 缺少 revision，无法验证分批读取一致性"
                )
            if revision is not None and current != revision:
                raise CloudAPIError(
                    "changed_during_read", "Sheets revision 在分批读取期间变化，请重试"
                )
            revision = current
            value_range = data.get("valueRange")
            if not isinstance(value_range, dict) or value_range.get("range") != area:
                raise CloudAPIError("invalid_response", "Sheets 返回区域与请求不一致")
            values = value_range.get("values", [])
            if not isinstance(values, list) or len(values) > last - first + 1:
                raise CloudAPIError("invalid_response", "Sheets 返回行数错误")
            for r in range(first, last + 1):
                cells = values[r - first] if r - first < len(values) else []
                if not isinstance(cells, list) or len(cells) > right - left + 1:
                    raise CloudAPIError("invalid_response", "Sheets 返回列数错误")
                row = []
                for c in range(left, right + 1):
                    value = cells[c - left] if c - left < len(cells) else None
                    if value is not None and not isinstance(
                        value, (str, int, float, bool)
                    ):
                        raise CloudAPIError(
                            "unsupported",
                            "Sheets 复杂单元格不可解释，请缩小范围或转换为静态文本",
                        )
                    value = str(value) if value is not None else None
                    self.limits.check(
                        len(value or ""), self.limits.max_cell_chars, "单元格字符数"
                    )
                    row.append(TableCell(r, c, value))
                rows.append(TableRow(r, tuple(row)))
        title = sheet.get("title") or sid
        table = build_table(
            source="feishu_sheets",
            identity=f"{token}:{sid}",
            title=title,
            rows=tuple(rows),
            header_rows=selection["header_rows"],
            merges=tuple(merges),
            region=selection["range"],
            section_path=(title,),
        )
        return self._table_parts(table), [table], [], revision, [], [], title

    def _base(self, resource, selection):
        token, tid = resource.token, selection["table_id"]
        prefix = f"/open-apis/bitable/v1/apps/{token}"
        app = self.client.get(prefix).get("app", {})
        fields = self.client.pages(
            prefix + f"/tables/{tid}/fields", params={"page_size": 100}
        )
        by_id = {f["field_id"]: f for f in fields}
        if len(by_id) != len(fields) or any(
            i not in by_id for i in selection["field_ids"]
        ):
            raise CloudAPIError("scope_changed", "选定字段已删除/重复，请重新检查范围")
        selected = [by_id[i] for i in selection["field_ids"]]
        self.limits.check(len(selected), self.limits.max_columns, "Base 选定字段数")
        names = [f["field_name"] for f in selected]
        if len(set(names)) != len(names):
            raise CloudAPIError(
                "invalid_response", "Base 字段名重复，无法可靠关联记录值"
            )
        body = {"field_names": names, "automatic_fields": True}
        for key in ("view_id", "filter"):
            if key in selection:
                body[key] = selection[key]
        records = self.client.pages(
            prefix + f"/tables/{tid}/records/search",
            params={"page_size": 500, "user_id_type": "open_id"},
            body=body,
        )
        self.limits.check(len(records), self.limits.max_rows, "Base 记录数")
        self.limits.check(
            len(records) * len(selected), self.limits.max_cells, "Base 单元格数"
        )
        missing = [
            f"{tid}/{f['field_id']}（{f['field_name']}）：暂不支持字段类型 {f['type']}"
            for f in selected
            if f["type"] not in SUPPORTED_FIELD_TYPES
        ]
        rows, seen = [], set()
        zone = app.get("time_zone") or "UTC"
        for ordinal, record in enumerate(records, 1):
            rid = resource_id(record.get("record_id"))
            if rid in seen:
                raise CloudAPIError("pagination", "Base 记录跨页重复，未发布快照")
            seen.add(rid)
            cells = []
            for c, field in enumerate(selected, 1):
                try:
                    value = field_text(
                        field,
                        (record.get("fields") or {}).get(field["field_name"]),
                        timezone=zone,
                    )
                except (FieldValueError, KeyError, TypeError) as exc:
                    value = f"[未提取：{exc}]"
                    missing.append(
                        f"{tid}/{rid}/{field['field_id']}（{field['field_name']}，类型 {field['type']}）：{exc}"
                    )
                self.limits.check(
                    len(value or ""), self.limits.max_cell_chars, "Base 字段字符数"
                )
                cells.append(TableCell(ordinal, c, value))
            rows.append(TableRow(ordinal, tuple(cells), rid))
        columns = tuple(
            TableColumn(c, (f["field_name"],), f["field_name"] + f" [{f['field_id']}]")
            for c, f in enumerate(selected, 1)
        )
        title = f"{app.get('name') or token} / {tid}"
        table = StructuredTable(
            "feishu_base", f"{token}:{tid}", title, columns, tuple(rows), region=tid
        )
        # Base has no snapshot revision in the search response. Record the
        # returned app revision separately when present, without claiming atomic remote reads.
        pieces = self._table_parts(table)
        return pieces, [table], missing, app.get("revision"), [], [], title

    def _docx(self, resource, selection):
        token = resource.token
        prefix = f"/open-apis/docx/v1/documents/{token}"
        document = self.client.get(prefix).get("document", {})
        revision = document.get("revision_id")
        if revision is None:
            raise CloudAPIError("invalid_response", "docx 元数据缺少 revision_id")
        blocks = self.client.pages(
            prefix + "/blocks",
            # Only -1 is the explicit latest-version read contract. Historical
            # numbered reads can require edit permission. Recheck metadata after
            # all pages instead of escalating a read-only app's scopes.
            params={"page_size": 500, "document_revision_id": -1},
        )
        by_id = {b["block_id"]: b for b in blocks}
        if len(by_id) != len(blocks):
            raise CloudAPIError("pagination", "docx block_id 重复")
        root = by_id.get(token) or next(
            (b for b in blocks if b.get("block_type") == 1), None
        )
        if root is None:
            raise CloudAPIError("invalid_response", "docx 缺少文档根块")
        title = document.get("title") or token
        pieces, tables, missing, children, references = [], [], [], [], []
        visited, headings = set(), []

        def block(identity):
            if identity not in by_id:
                raise CloudAPIError("pagination", f"docx 缺少引用块 {identity}")
            if identity in visited:
                raise CloudAPIError("invalid_response", "docx 块循环或重复挂载")
            visited.add(identity)
            return by_id[identity]

        def text_of(item):
            field = next(
                (
                    k
                    for k in (
                        "page",
                        "text",
                        *(f"heading{i}" for i in range(1, 10)),
                        "bullet",
                        "ordered",
                        "code",
                        "quote",
                        "todo",
                    )
                    if k in item
                ),
                None,
            )
            if not field:
                return ""
            result = []
            for element in item[field].get("elements", []):
                if "text_run" in element:
                    run = element["text_run"]
                    value = run.get("content", "")
                    link = (
                        (run.get("text_element_style") or {}).get("link", {}).get("url")
                    )
                    result.append(value + (f" ({link})" if link else ""))
                elif "equation" in element:
                    result.append(element["equation"].get("content", ""))
                elif "mention_user" in element:
                    result.append(
                        "[人员引用："
                        + str(element["mention_user"].get("user_id", "未知"))
                        + "]"
                    )
                elif "mention_doc" in element:
                    ref = element["mention_doc"]
                    result.append(
                        f"{ref.get('title') or '文档引用'} ({ref.get('url') or ref.get('token') or '未知'})"
                    )
                else:
                    missing.append(f"docx block {item['block_id']}：未支持的富文本元素")
            return "".join(result)

        def cell_text(identity, depth=0):
            if depth > 64:
                raise CloudAPIError("limit", "docx 单元格嵌套过深")
            item = block(identity)
            if item.get("block_type") in {18, 30, 31}:
                raise ValueError("单元格内嵌复杂表格/引用暂不可恢复")
            values = [text_of(item)]
            values.extend(
                cell_text(child, depth + 1) for child in item.get("children", [])
            )
            return "\n".join(v for v in values if v)

        def walk(identity, depth=0):
            nonlocal headings
            if depth > 64:
                raise CloudAPIError("limit", "docx 块嵌套过深")
            item = block(identity)
            kind = item.get("block_type")
            locator = f"feishu_docx:{token}:block:{identity}"
            if kind == 31:
                spec = item.get("table", {})
                prop = spec.get("property", {})
                height, width = prop.get("row_size", 0), prop.get("column_size", 0)
                self.limits.check(height, self.limits.max_rows, "docx 表格行数")
                self.limits.check(width, self.limits.max_columns, "docx 表格列数")
                self.limits.check(
                    height * width, self.limits.max_cells, "docx 表格单元格数"
                )
                cells = spec.get("cells", [])
                if height < 1 or width < 1 or len(cells) != height * width:
                    raise CloudAPIError(
                        "invalid_response", "docx 表格维度与单元格不匹配"
                    )
                rows, merges = [], []
                info = prop.get("merge_info", [])
                if info and len(info) != len(cells):
                    raise CloudAPIError(
                        "invalid_response", "docx 合并关系数与表格不一致"
                    )
                covered = set()
                for r in range(1, height + 1):
                    values = []
                    for c in range(1, width + 1):
                        index = (r - 1) * width + c - 1
                        value = cell_text(cells[index])
                        self.limits.check(
                            len(value), self.limits.max_cell_chars, "docx 单元格字符数"
                        )
                        values.append(TableCell(r, c, value or None))
                        if info and (r, c) not in covered:
                            rs, cs = (
                                info[index].get("row_span", 1),
                                info[index].get("col_span", 1),
                            )
                            if (
                                min(rs, cs) < 1
                                or r + rs - 1 > height
                                or c + cs - 1 > width
                            ):
                                raise CloudAPIError(
                                    "invalid_response", "docx 合并区域越界"
                                )
                            if rs > 1 or cs > 1:
                                merges.append(TableMerge(r, c, r + rs - 1, c + cs - 1))
                                covered.update(
                                    (rr, cc)
                                    for rr in range(r, r + rs)
                                    for cc in range(c, c + cs)
                                )
                    rows.append(TableRow(r, tuple(values)))
                table = build_table(
                    source="feishu_docx",
                    identity=f"{token}:block:{identity}",
                    title=f"{title} / 表格 {identity}",
                    rows=tuple(rows),
                    header_rows=selection["header_rows"],
                    merges=tuple(merges),
                    region=identity,
                    section_path=tuple(headings),
                )
                tables.append(table)
                pieces.extend(self._table_parts(table))
                return
            if kind in {18, 30}:
                ref_kind = "base" if kind == 18 else "sheets"
                raw = (item.get("bitable" if kind == 18 else "sheet") or {}).get(
                    "token", ""
                )
                references.append(
                    {"block_id": identity, "kind": ref_kind, "reference_token": raw}
                )
                explicit = selection["embeds"].get(identity)
                if not explicit:
                    missing.append(
                        f"{locator}：内嵌 {ref_kind} 未选择范围（引用 {raw}），未读取"
                    )
                    return
                if not isinstance(explicit, dict) or set(explicit) != {
                    "url",
                    "selection",
                }:
                    raise ValueError("embeds 配置必须含 url 和 selection")
                child_resource = self.resolve(explicit["url"])
                child_selection = normalize_selection(
                    child_resource, explicit["selection"]
                )
                secondary = child_selection.get("sheet_id") or child_selection.get(
                    "table_id"
                )
                if child_resource.kind != ref_kind or raw not in {
                    child_resource.token,
                    child_resource.token + "_" + secondary,
                }:
                    raise ValueError(
                        "显式内嵌选择与块引用身份不一致；不会猜测拆分 token"
                    )
                try:
                    child = self.load(child_resource, child_selection)
                    children.append(child)
                    references[-1].update(
                        source_ref=child.source_ref,
                        state="partial" if child.parsed.missing_content else "read",
                    )
                    self.limits.check(
                        sum(len(s.parsed.text) for s in children),
                        self.limits.max_output_chars,
                        "嵌入来源总输出字符数",
                    )
                    if child.parsed.missing_content:
                        missing.append(
                            f"{locator}：内嵌来源部分成功：{'; '.join(child.parsed.missing_content)}"
                        )
                except (CloudAPIError, ValueError) as exc:
                    missing.append(
                        f"{locator}：内嵌来源 {explicit['url']} 读取失败：{exc}"
                    )
                    references[-1].update(
                        failure_kind=exc.kind
                        if isinstance(exc, CloudAPIError)
                        else "unsupported",
                        reason=str(exc),
                        source_ref=self.source_ref(child_resource, child_selection),
                    )
                return
            value = text_of(item)
            if value:
                if isinstance(kind, int) and 3 <= kind <= 11:
                    level = kind - 2
                    headings = headings[: level - 1] + [value]
                if kind == 14:
                    value = "```\n" + value + "\n```"
                pieces.append((value, tuple(headings), locator, "text"))
            elif kind not in {1, 12, 13, 14, 15, 17, 19, 24, 25, 32, 34}:
                missing.append(f"{locator}：未提取块类型 {kind}")
            for child in item.get("children", []):
                walk(child, depth + 1)

        walk(root["block_id"])
        if set(by_id) - visited:
            raise CloudAPIError(
                "invalid_response", "docx 存在未关联到文档树的块，未发布不完整快照"
            )
        if set(selection["embeds"]) - {r["block_id"] for r in references}:
            raise CloudAPIError("scope_changed", "已选择的内嵌块已消失，请重新检查范围")
        if self.client.get(prefix).get("document", {}).get("revision_id") != revision:
            raise CloudAPIError("changed_during_read", "docx 在读取期间已变化，请重试")
        if not pieces:
            pieces.append(("[文档无可提取正文]", (), f"feishu_docx:{token}", "text"))
        return pieces, tables, missing, revision, children, references, title
